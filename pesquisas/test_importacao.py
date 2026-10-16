from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import NamedTemporaryFile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from .importacao import ErroImportacao, importar_colaboradores, preparar_linhas, validar_linhas
from .models import Colaborador, Empresa, ImportacaoColaboradores
from .test_utils import criar_usuario


CPF_ANA = '52998224725'
CPF_BRUNO = '11144477735'
CPF_ZERO = '09784735903'  # Começa com zero: planilhas costumam perdê-lo


def csv_bytes(texto, codificacao='utf-8'):
    return texto.encode(codificacao)


def xlsx_bytes(linhas):
    pasta = Workbook()
    for linha in linhas:
        pasta.active.append(linha)
    saida = BytesIO()
    pasta.save(saida)
    return saida.getvalue()


class ImportacaoServicoTests(TestCase):
    def test_importa_csv_so_com_colunas_obrigatorias(self):
        resultado = importar_colaboradores(csv_bytes(
            'nome;documento;empresa\n'
            f'Ana Silva;{CPF_ANA};Matriz\n'
            f'Bruno Souza;{CPF_BRUNO};matriz\n'
        ), 'colaboradores.csv')

        self.assertEqual((resultado.criados, resultado.atualizados), (2, 0))
        self.assertEqual(Empresa.objects.filter(nome__iexact='matriz').count(), 1)
        ana = Colaborador.objects.get(documento=CPF_ANA)
        self.assertTrue(ana.ativo)
        self.assertIsNotNone(ana.data_ativacao)
        self.assertEqual((ana.celulares, ana.email, ana.data_contratacao), ([], '', None))

    def test_importa_xlsx_com_todos_os_campos(self):
        resultado = importar_colaboradores(xlsx_bytes([
            ['Nome', 'CPF', 'Unidade Fabril', 'Celular', 'E-mail', 'Data de contratação', 'Data de demissão',
             'É gestor', 'Nome do gestor', 'Observação'],
            ['Bruno Souza', CPF_BRUNO, 'Matriz', 47999998888, 'Bruno@Exemplo.com.br', datetime(2024, 3, 1), None,
             'Não', 'ana  GESTORA', 'qualquer coisa'],
            ['Ana Gestora', int(CPF_ZERO), 'Filial', '(47) 98888-7777', '', '01/02/2020', '', 'Sim', '', ''],
        ]), 'colaboradores.xlsx')

        self.assertEqual(resultado.criados, 2)
        self.assertEqual(resultado.colunas_ignoradas, ['observacao'])
        bruno = Colaborador.objects.get(documento=CPF_BRUNO)
        ana = Colaborador.objects.get(documento=CPF_ZERO)
        self.assertEqual(bruno.celulares, ['(47) 99999-8888'])
        self.assertEqual(bruno.email, 'bruno@exemplo.com.br')
        self.assertEqual(bruno.data_contratacao, date(2024, 3, 1))
        self.assertFalse(bruno.eh_gestor)
        self.assertEqual(bruno.gestor, ana)  # gestora marcada mais abaixo na planilha
        self.assertTrue(ana.eh_gestor)
        self.assertEqual(ana.data_contratacao, date(2020, 2, 1))
        self.assertEqual(ana.nomes_empresas, 'Filial')

    def test_gestor_pelo_nome_usa_gestores_ativos_ja_cadastrados(self):
        matriz = Empresa.objects.create(nome='Matriz')
        joao = Colaborador.objects.create(nome='João Conceição', documento=CPF_ZERO, empresas=[matriz], eh_gestor=True)
        importar_colaboradores(csv_bytes(f'nome;documento;empresa;gestor\nAna;{CPF_ANA};Matriz;joao conceicao\n'), 'a.csv')
        self.assertEqual(Colaborador.objects.get(documento=CPF_ANA).gestor, joao)

        # Célula vazia mantém o gestor atual.
        importar_colaboradores(csv_bytes(f'nome;documento;empresa;gestor\nAna;{CPF_ANA};Matriz;\n'), 'a.csv')
        self.assertEqual(Colaborador.objects.get(documento=CPF_ANA).gestor, joao)

    def test_gestor_pelo_nome_aponta_problemas(self):
        matriz = Empresa.objects.create(nome='Matriz')
        Colaborador.objects.create(nome='Comum', documento=CPF_ZERO, empresas=[matriz])
        Colaborador.objects.create(nome='Inativo', documento='12345678909', empresas=[matriz], eh_gestor=True, ativo=False)
        Colaborador.objects.create(nome='Dupla', documento='12345678062', empresas=[matriz], eh_gestor=True)
        Colaborador.objects.create(nome='Dupla', documento='98765432029', empresas=[matriz], eh_gestor=True)

        with self.assertRaises(ErroImportacao) as ctx:
            importar_colaboradores(csv_bytes(
                'nome;documento;empresa;eh_gestor;gestor\n'
                f'Ana;{CPF_ANA};Matriz;;Comum\n'
                f'Bruno;{CPF_BRUNO};Matriz;Sim;Bruno\n'
                f'Ana;{CPF_ANA};Matriz;;\n'
                f'Carla;11122233396;Matriz;;Inativo\n'
                f'Davi;39053344705;Matriz;;Dupla\n'
                f'Eva;71428793860;Matriz;;Ninguém\n'
            ), 'a.csv')
        erros = '\n'.join(ctx.exception.erros)
        self.assertIn('Linha 2: Gestor: "Comum" não está marcado como gestor ativo', erros)
        self.assertIn('Linha 3: Gestor: O colaborador não pode ser gestor de si mesmo', erros)
        self.assertIn('Linha 4: Este CPF/CNPJ tem erro em outra linha (2)', erros)
        self.assertIn('Linha 5: Gestor: "Inativo" não está marcado como gestor ativo', erros)
        self.assertIn('Há 2 gestores chamados "Dupla"', erros)
        self.assertIn('Nenhum gestor chamado "Ninguém"', erros)
        self.assertFalse(Colaborador.objects.filter(documento=CPF_ANA).exists())

    def test_pode_desmarcar_gestor_se_os_liderados_trocam_de_gestor_na_planilha(self):
        matriz = Empresa.objects.create(nome='Matriz')
        antiga = Colaborador.objects.create(nome='Antiga', documento=CPF_ZERO, empresas=[matriz], eh_gestor=True)
        Colaborador.objects.create(nome='Ana', documento=CPF_ANA, empresas=[matriz], gestor=antiga)

        with self.assertRaisesMessage(ErroImportacao, 'não pode deixar de ser gestor'):
            importar_colaboradores(csv_bytes(f'nome;documento;empresa;eh_gestor\nAntiga;{CPF_ZERO};Matriz;Não\n'), 'a.csv')

        importar_colaboradores(csv_bytes(
            'nome;documento;empresa;eh_gestor;gestor\n'
            f'Antiga;{CPF_ZERO};Matriz;Não;\n'
            f'Nova;{CPF_BRUNO};Matriz;Sim;\n'
            f'Ana;{CPF_ANA};Matriz;;Nova\n'
        ), 'a.csv')
        self.assertEqual(Colaborador.objects.get(documento=CPF_ANA).gestor.nome, 'Nova')
        self.assertFalse(Colaborador.objects.get(documento=CPF_ZERO).eh_gestor)

    def test_atualiza_existente_sem_apagar_campos_vazios_nem_mudar_status(self):
        empresa = Empresa.objects.create(nome='Matriz')
        Colaborador.objects.create(
            nome='Ana', documento=CPF_ANA, empresas=[empresa], ativo=False,
            celulares=['47999998888'], email='ana@exemplo.com.br', data_contratacao=date(2020, 1, 1),
        )
        resultado = importar_colaboradores(csv_bytes(
            'nome,documento,empresa,celular,email,data_contratacao\n'
            f'Ana Souza,{CPF_ANA},Matriz,,,\n'
        ), 'a.csv')

        self.assertEqual((resultado.criados, resultado.atualizados), (0, 1))
        ana = Colaborador.objects.get(documento=CPF_ANA)
        self.assertEqual(ana.nome, 'Ana Souza')
        self.assertFalse(ana.ativo)
        self.assertEqual((ana.celulares, ana.email, ana.data_contratacao), (['(47) 99999-8888'], 'ana@exemplo.com.br', date(2020, 1, 1)))

    def test_importa_data_de_nascimento(self):
        importar_colaboradores(csv_bytes(
            'nome;documento;empresa;Data de nascimento\n'
            f'Ana;{CPF_ANA};Matriz;20/05/1990\n'
            f'Bruno;{CPF_BRUNO};Matriz;\n'
        ), 'a.csv')
        self.assertEqual(Colaborador.objects.get(documento=CPF_ANA).data_nascimento, date(1990, 5, 20))
        self.assertIsNone(Colaborador.objects.get(documento=CPF_BRUNO).data_nascimento)

        # Célula vazia mantém a data atual.
        importar_colaboradores(csv_bytes(f'nome;documento;empresa;data_nascimento\nAna;{CPF_ANA};Matriz;\n'), 'a.csv')
        self.assertEqual(Colaborador.objects.get(documento=CPF_ANA).data_nascimento, date(1990, 5, 20))

    def test_data_de_nascimento_futura_e_recusada(self):
        with self.assertRaises(ErroImportacao) as erro:
            importar_colaboradores(csv_bytes(f'nome;documento;empresa;data_nascimento\nAna;{CPF_ANA};Matriz;01/01/2999\n'), 'a.csv')
        self.assertIn('Linha 2', erro.exception.erros[0])
        self.assertIn('nascimento', erro.exception.erros[0])
        self.assertFalse(Colaborador.objects.exists())

    def test_mesmo_cpf_em_varias_linhas_vira_um_colaborador_em_varias_empresas(self):
        resultado = importar_colaboradores(csv_bytes(
            'nome;documento;empresa;email\n'
            f'LEODATO JUNKES;{CPF_ANA};1 - Matriz;\n'
            f'Leodato  Junkes;{CPF_ANA};20 - EVO;leodato@exemplo.com.br\n'
            f'LEODATO JUNKES;{CPF_ANA};2 - Filial MG;\n'
            f'LEODATO JUNKES;{CPF_ANA};1 - matriz;\n'
        ), 'a.csv')

        self.assertEqual((resultado.criados, resultado.atualizados), (1, 0))
        leodato = Colaborador.objects.get()
        self.assertEqual(leodato.nomes_empresas, '1 - Matriz, 2 - Filial MG, 20 - EVO')
        self.assertEqual(leodato.email, 'leodato@exemplo.com.br')

    def test_varios_celulares_na_celula_e_em_linhas_diferentes(self):
        importar_colaboradores(xlsx_bytes([
            ['nome', 'documento', 'empresa', 'celulares'],
            ['Ana', CPF_ANA, 'Matriz', '47999998888 / (11) 98888-7777'],
            ['Ana', CPF_ANA, 'Filial', 21977776666],
            ['Ana', CPF_ANA, 'EVO', '(47) 99999-8888'],
        ]), 'a.xlsx')
        self.assertEqual(Colaborador.objects.get().celulares, ['(47) 99999-8888', '(11) 98888-7777', '(21) 97777-6666'])

    def test_celulares_da_planilha_sao_somados_aos_ja_cadastrados(self):
        matriz = Empresa.objects.create(nome='Matriz')
        Colaborador.objects.create(nome='Ana', documento=CPF_ANA, empresas=[matriz], celulares=['47999998888'])
        importar_colaboradores(csv_bytes(f'nome;documento;empresa;celular\nAna;{CPF_ANA};Matriz;11988887777, 47999998888\n'), 'a.csv')
        self.assertEqual(Colaborador.objects.get().celulares, ['(47) 99999-8888', '(11) 98888-7777'])

    def test_data_de_demissao_na_planilha_inativa(self):
        matriz = Empresa.objects.create(nome='Matriz')
        Colaborador.objects.create(nome='Ana', documento=CPF_ANA, empresas=[matriz])
        importar_colaboradores(csv_bytes(
            f'nome;documento;empresa;data_demissao\nAna;{CPF_ANA};Matriz;30/09/2026\nBruno;{CPF_BRUNO};Matriz;30/09/2026\n'
        ), 'a.csv')
        for documento in (CPF_ANA, CPF_BRUNO):
            colaborador = Colaborador.objects.get(documento=documento)
            self.assertFalse(colaborador.ativo)
            self.assertEqual(colaborador.data_demissao, date(2026, 9, 30))
            self.assertIsNotNone(colaborador.data_inativacao)

    def test_reimportacao_define_as_empresas_do_colaborador(self):
        matriz = Empresa.objects.create(nome='Matriz')
        Colaborador.objects.create(nome='Ana', documento=CPF_ANA, empresas=[matriz])
        importar_colaboradores(csv_bytes(f'nome;documento;empresa\nAna;{CPF_ANA};EVO\nAna;{CPF_ANA};Filial\n'), 'a.csv')
        self.assertEqual(Colaborador.objects.get().nomes_empresas, 'EVO, Filial')

    def test_qualquer_erro_cancela_toda_a_importacao_e_lista_todas_as_linhas(self):
        with self.assertRaises(ErroImportacao) as ctx:
            importar_colaboradores(csv_bytes(
                'nome;documento;empresa;celular;email;data_contratacao;eh_gestor\n'
                f'Ana;{CPF_ANA};Matriz;;;;\n'
                'Sem CPF;;Matriz;;;;\n'
                f'Bruno;{CPF_BRUNO};Matriz;4733334444;bruno@;31/02/2024;\n'
                f'Repetida;{CPF_ANA};Matriz;;;;\n'
                f'Carla;{CPF_ZERO};Matriz;;;;talvez\n'
            ), 'a.csv')

        erros = ctx.exception.erros
        self.assertTrue(any(e.startswith('Linha 2:') and 'erro em outra linha (5)' in e for e in erros), erros)
        self.assertTrue(any(e.startswith('Linha 3:') and 'CPF/CNPJ: Preenchimento obrigatório' in e for e in erros), erros)
        self.assertTrue(any(e.startswith('Linha 4:') and 'Contratação: Data "31/02/2024" inválida' in e for e in erros), erros)
        self.assertTrue(any(e.startswith('Linha 5:') and 'outro nome' in e for e in erros), erros)
        self.assertTrue(any(e.startswith('Linha 6:') and 'É gestor' in e for e in erros), erros)
        self.assertFalse(Colaborador.objects.exists())
        self.assertFalse(Empresa.objects.exists())

    def test_celular_e_email_invalidos_sao_apontados(self):
        with self.assertRaises(ErroImportacao) as ctx:
            importar_colaboradores(csv_bytes(
                f'nome,documento,empresa,celular,email\nBruno,{CPF_BRUNO},Matriz,4733334444,bruno@\n'
            ), 'a.csv')
        erro = ctx.exception.erros[0]
        self.assertIn('Celular: Celular inválido', erro)
        self.assertIn('E-mail: Informe um e-mail válido', erro)

    def test_csv_do_excel_em_windows_1252(self):
        importar_colaboradores(csv_bytes(f'nome;documento;empresa\nJoão Conceição;{CPF_ANA};Matriz\n', 'cp1252'), 'a.csv')
        self.assertEqual(Colaborador.objects.get().nome, 'João Conceição')

    def test_rejeita_coluna_obrigatoria_ausente_e_extensao_invalida(self):
        with self.assertRaisesMessage(ErroImportacao, 'empresa'):
            importar_colaboradores(csv_bytes(f'nome,documento\nAna,{CPF_ANA}\n'), 'a.csv')
        with self.assertRaisesMessage(ErroImportacao, '.csv ou .xlsx'):
            importar_colaboradores(b'qualquer', 'a.xls')


class ValidacaoLinhasTests(TestCase):
    def linhas(self, texto):
        return preparar_linhas(csv_bytes(texto), 'a.csv')[0]

    def test_erros_apontam_a_celula(self):
        erros = validar_linhas(self.linhas(
            'nome;documento;empresa;email;data_nascimento\n'
            f'Ana;{CPF_ANA};Matriz;ana@exemplo.com.br;20/05/1990\n'
            f'Bruno;123;;bruno@;31/02/1990\n'
        ))
        self.assertEqual(list(erros), [3])
        self.assertEqual(set(erros[3]), {'documento', 'empresa', 'email', 'data_nascimento'})

    def test_gestor_em_linha_com_erro_segura_o_liderado(self):
        linhas = self.linhas(
            'nome;documento;empresa;eh_gestor;gestor\n'
            f'Gestora;123;Matriz;Sim;\n'
            f'Ana;{CPF_ANA};Matriz;;Gestora\n'
        )
        erros = validar_linhas(linhas)
        self.assertIn('linha 2, que tem erro', erros[3]['gestor'])

        linhas[0]['valores']['documento'] = CPF_BRUNO  # corrigida a gestora, o liderado também fica válido
        self.assertEqual(validar_linhas(linhas), {})

    def test_xlsx_vira_texto_editavel(self):
        linhas, _ = preparar_linhas(xlsx_bytes([
            ['nome', 'documento', 'empresa', 'celular', 'data_contratacao', 'eh_gestor'],
            ['Ana', int(CPF_ZERO), 'Matriz', 47999998888, datetime(2024, 3, 1), True],
        ]), 'a.xlsx')
        self.assertEqual(linhas[0]['numero'], 2)
        valores = linhas[0]['valores']
        self.assertEqual((valores['documento'], valores['celular'], valores['data_contratacao'], valores['eh_gestor']),
                         (CPF_ZERO, '47999998888', '01/03/2024', 'Sim'))
        self.assertEqual(valores['email'], '')


class ImportacaoComandoTests(TestCase):
    def test_comando_completo_importa_e_falha_sem_gravar(self):
        with NamedTemporaryFile('w', encoding='utf-8', suffix='.csv', delete=False, newline='') as arquivo:
            arquivo.write(f'nome;documento;empresa;email\nAna;{CPF_ANA};Matriz;ana@exemplo.com.br\n')
            caminho = Path(arquivo.name)
        try:
            saida = StringIO()
            call_command('importar_colaboradores_completo', caminho, stdout=saida)
            self.assertIn('1 criados', saida.getvalue())
            self.assertEqual(Colaborador.objects.get().email, 'ana@exemplo.com.br')

            caminho.write_text('nome;documento;empresa\nSem CPF;;Matriz\n', encoding='utf-8')
            with self.assertRaisesMessage(CommandError, 'Linha 2'):
                call_command('importar_colaboradores_completo', caminho, stdout=StringIO())
        finally:
            caminho.unlink(missing_ok=True)


class ImportacaoTelaTests(TestCase):
    def setUp(self):
        self.client.force_login(criar_usuario('geral', modulos=['geral']))
        self.url = reverse('pesquisas:colaborador_importar')

    def test_lista_tem_botao_de_importar(self):
        self.assertContains(self.client.get(reverse('pesquisas:colaborador_lista')), self.url)

    def enviar(self, conteudo, nome='a.csv'):
        return self.client.post(self.url, {'arquivo': SimpleUploadedFile(nome, conteudo)})

    def test_upload_leva_para_a_revisao_sem_gravar_e_a_confirmacao_importa(self):
        resposta = self.enviar(xlsx_bytes([
            ['nome', 'documento', 'empresa'], ['Ana', CPF_ANA, 'Matriz'], ['Bruno', CPF_BRUNO, 'Matriz'],
        ]), 'colaboradores.xlsx')
        importacao = ImportacaoColaboradores.objects.get()
        self.assertRedirects(resposta, reverse('pesquisas:colaborador_importar_revisao', args=[importacao.pk]))
        self.assertFalse(Colaborador.objects.exists())

        revisao = self.client.get(resposta.url)
        self.assertContains(revisao, 'id="qtdValidos">2<')
        self.assertContains(revisao, 'id="qtdErros">0<')

        resposta = self.client.post(reverse('pesquisas:colaborador_importar_confirmar', args=[importacao.pk]), follow=True)
        self.assertRedirects(resposta, reverse('pesquisas:colaborador_lista'))
        self.assertContains(resposta, '2 criado(s)')
        self.assertEqual(Colaborador.objects.count(), 2)
        self.assertFalse(ImportacaoColaboradores.objects.exists())

    def test_arquivo_ilegivel_continua_no_upload(self):
        resposta = self.enviar(csv_bytes(f'nome,documento\nAna,{CPF_ANA}\n'))
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, 'Não foi possível ler a planilha')
        self.assertFalse(ImportacaoColaboradores.objects.exists())

    def test_corrige_celula_e_a_linha_vira_valida(self):
        self.enviar(csv_bytes(f'nome,documento,empresa,email\nAna,{CPF_ANA},Matriz,\nBruno,111,Matriz,bruno@\n'))
        importacao = ImportacaoColaboradores.objects.get()
        url = reverse('pesquisas:colaborador_importar_celula', args=[importacao.pk])
        editar = lambda coluna, valor: self.client.post(  # noqa: E731
            url, {'numero': 3, 'coluna': coluna, 'valor': valor}, content_type='application/json').json()

        estado = editar('documento', '111.444.777-35')
        self.assertEqual((estado['total'], estado['com_erro']), (2, 1))
        self.assertEqual(list(estado['erros']['3']), ['email'])  # o erro do CPF sumiu; o do e-mail continua

        estado = editar('email', '  bruno@exemplo.com.br ')
        self.assertEqual((estado['com_erro'], estado['erros'], estado['valor']), (0, {}, 'bruno@exemplo.com.br'))
        self.assertFalse(Colaborador.objects.exists())  # nada é gravado antes da confirmação

        self.client.post(reverse('pesquisas:colaborador_importar_confirmar', args=[importacao.pk]))
        self.assertEqual(Colaborador.objects.get(documento=CPF_BRUNO).email, 'bruno@exemplo.com.br')

    def test_confirmar_grava_so_as_validas(self):
        self.enviar(csv_bytes(f'nome,documento,empresa\nAna,{CPF_ANA},Matriz\nSem CPF,,Matriz\n'))
        importacao = ImportacaoColaboradores.objects.get()
        resposta = self.client.post(reverse('pesquisas:colaborador_importar_confirmar', args=[importacao.pk]), follow=True)
        self.assertContains(resposta, '1 criado(s)')
        self.assertContains(resposta, '1 linha(s) com erro não foram importadas')
        self.assertEqual(list(Colaborador.objects.values_list('nome', flat=True)), ['Ana'])

    def test_sem_linhas_validas_nao_confirma(self):
        self.enviar(csv_bytes('nome,documento,empresa\nSem CPF,,Matriz\n'))
        importacao = ImportacaoColaboradores.objects.get()
        resposta = self.client.post(reverse('pesquisas:colaborador_importar_confirmar', args=[importacao.pk]), follow=True)
        self.assertContains(resposta, 'Não há linhas válidas para importar')
        self.assertTrue(ImportacaoColaboradores.objects.exists())

    def test_descarta_linhas_com_erro_em_lote_e_uma_a_uma(self):
        self.enviar(csv_bytes(
            f'nome,documento,empresa\nAna,{CPF_ANA},Matriz\nSem CPF,,Matriz\nSem empresa,{CPF_BRUNO},\nCarla,{CPF_ZERO},Matriz\n'
        ))
        importacao = ImportacaoColaboradores.objects.get()
        url = reverse('pesquisas:colaborador_importar_descartar', args=[importacao.pk])

        estado = self.client.post(url, {'numeros': [5]}, content_type='application/json').json()
        self.assertEqual((estado['total'], estado['com_erro'], estado['descartadas']), (3, 2, [5]))

        estado = self.client.post(url, {'com_erro': True}, content_type='application/json').json()
        self.assertEqual((estado['total'], estado['com_erro'], estado['descartadas']), (1, 0, [3, 4]))
        importacao.refresh_from_db()
        self.assertEqual([linha['numero'] for linha in importacao.linhas], [2])

    def test_cancelar_descarta_a_revisao(self):
        self.enviar(csv_bytes(f'nome,documento,empresa\nAna,{CPF_ANA},Matriz\n'))
        importacao = ImportacaoColaboradores.objects.get()
        self.client.post(reverse('pesquisas:colaborador_importar_cancelar', args=[importacao.pk]))
        self.assertFalse(ImportacaoColaboradores.objects.exists())
        self.assertFalse(Colaborador.objects.exists())

    def test_revisao_de_outro_usuario_nao_e_acessivel(self):
        self.enviar(csv_bytes(f'nome,documento,empresa\nAna,{CPF_ANA},Matriz\n'))
        importacao = ImportacaoColaboradores.objects.get()
        self.client.force_login(criar_usuario('outro', modulos=['geral']))
        self.assertEqual(self.client.get(reverse('pesquisas:colaborador_importar_revisao', args=[importacao.pk])).status_code, 404)
        resposta = self.client.post(reverse('pesquisas:colaborador_importar_confirmar', args=[importacao.pk]))
        self.assertEqual(resposta.status_code, 404)
        self.assertFalse(Colaborador.objects.exists())

    def test_edicao_rejeita_coluna_desconhecida(self):
        self.enviar(csv_bytes(f'nome,documento,empresa\nAna,{CPF_ANA},Matriz\n'))
        importacao = ImportacaoColaboradores.objects.get()
        resposta = self.client.post(reverse('pesquisas:colaborador_importar_celula', args=[importacao.pk]),
                                    {'numero': 2, 'coluna': 'ativo', 'valor': 'x'}, content_type='application/json')
        self.assertEqual(resposta.status_code, 400)

    def test_baixa_planilha_modelo(self):
        resposta = self.client.get(reverse('pesquisas:colaborador_modelo_importacao'))
        self.assertEqual(resposta['Content-Type'], 'text/csv; charset=utf-8')
        self.assertIn('nome;documento;empresa', resposta.content.decode('utf-8-sig'))

    def test_modelo_baixado_importa_sem_erros(self):
        modelo = self.client.get(reverse('pesquisas:colaborador_modelo_importacao')).content
        resultado = importar_colaboradores(modelo, 'modelo.csv')
        self.assertEqual(resultado.criados, 2)
        self.assertTrue(Colaborador.objects.get(documento=CPF_ANA).eh_gestor)
        self.assertEqual(Colaborador.objects.get(documento=CPF_BRUNO).gestor.documento, CPF_ANA)

    def test_sem_permissao_nao_acessa(self):
        self.client.force_login(criar_usuario('sem.modulo', modulos=[]))
        self.assertNotEqual(self.client.get(self.url).status_code, 200)


class ExportacaoColaboradoresTests(TestCase):
    def setUp(self):
        from openpyxl import load_workbook  # noqa: F401  (garante a dependência)

        self.matriz = Empresa.objects.create(nome='Matriz')
        self.filial = Empresa.objects.create(nome='Filial')
        self.gestora = Colaborador.objects.create(
            nome='Ana Gestora', documento=CPF_ANA, empresas=[self.matriz, self.filial], eh_gestor=True,
            celulares=['47999998888'], email='ana@exemplo.com.br', data_contratacao=date(2024, 3, 1),
            data_nascimento=date(1985, 5, 20),
        )
        Colaborador.objects.create(nome='Bruno', documento=CPF_BRUNO, empresas=[self.filial], gestor=self.gestora)
        Colaborador.objects.create(nome='Carla', documento=CPF_ZERO, empresas=[self.matriz], ativo=False)
        self.client.force_login(criar_usuario('admin.export', administrador=True))
        self.url = reverse('pesquisas:colaborador_exportar')

    def _planilha(self, resposta):
        from openpyxl import load_workbook

        return [list(linha) for linha in load_workbook(BytesIO(resposta.content)).active.iter_rows(values_only=True)]

    def test_exporta_xlsx_com_todos_os_dados(self):
        resposta = self.client.get(self.url)

        self.assertEqual(resposta.status_code, 200)
        self.assertIn('.xlsx', resposta['Content-Disposition'])
        linhas = self._planilha(resposta)
        self.assertEqual(linhas[0][:3], ['Nome', 'CPF / CNPJ', 'Empresas'])
        self.assertEqual([linha[0] for linha in linhas[1:]], ['Ana Gestora', 'Bruno', 'Carla'])
        ana, bruno, carla = linhas[1:]
        self.assertEqual(ana[2], 'Filial, Matriz')
        self.assertEqual(ana[4], '(47) 99999-8888')
        self.assertEqual(ana[6], datetime(1985, 5, 20))
        self.assertEqual(ana[7], datetime(2024, 3, 1))
        self.assertEqual(ana[9], 'Sim')
        self.assertEqual(bruno[10], 'Ana Gestora')
        self.assertEqual(carla[1], CPF_ZERO)  # zero à esquerda preservado
        self.assertEqual(carla[3], 'Inativo')

    def test_respeita_filtros_da_listagem(self):
        linhas = self._planilha(self.client.get(self.url, {'status': 'ativo', 'empresa': self.filial.pk}))
        self.assertEqual([linha[0] for linha in linhas[1:]], ['Ana Gestora', 'Bruno'])

        linhas = self._planilha(self.client.get(self.url, {'q': 'carla'}))
        self.assertEqual([linha[0] for linha in linhas[1:]], ['Carla'])

    def test_exporta_csv(self):
        resposta = self.client.get(self.url, {'formato': 'csv'})

        self.assertIn('.csv', resposta['Content-Disposition'])
        texto = resposta.content.decode('utf-8-sig')
        self.assertTrue(texto.startswith('Nome;CPF / CNPJ;Empresas;Status'))
        self.assertIn(f'Ana Gestora;{CPF_ANA};Filial, Matriz;Ativo;(47) 99999-8888;ana@exemplo.com.br;20/05/1985;01/03/2024;;Sim;', texto)

    def test_exige_login(self):
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)
